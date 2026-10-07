"""Tests de contrato para ``docling_explorer.chunker``."""

from pathlib import Path

import docx
import pytest
from docling.document_converter import DocumentConverter
from docling.exceptions import ConversionError
from docling_core.transforms.chunker import BaseChunk, HierarchicalChunker
from docling_core.types.doc import DoclingDocument as DLDocument
from pydantic import ValidationError
from transformers import AutoTokenizer

from docling_explorer.chunker import MaxTokenLimitingChunker, chunk_document

SAMPLE_PDF = Path(__file__).parents[1] / "resources" / "JerarquiaDocs.pdf"
SAMPLE_DOCX = Path(__file__).parents[1] / "resources" / "rel18" / "rel_14.docx"
MODEL_NAME = "BAAI/bge-small-en-v1.5"


@pytest.fixture(scope="session")
def tokenizer():
    return AutoTokenizer.from_pretrained(MODEL_NAME)


@pytest.fixture(scope="session")
def sample_document():
    converter = DocumentConverter()
    result = converter.convert(str(SAMPLE_PDF))
    return result.document


@pytest.fixture(scope="session")
def docx_with_headings(tmp_path_factory):
    """DOCX sintético con estilos de título reconocidos por Docling."""
    tmp = tmp_path_factory.mktemp("fixtures")
    path = tmp / "headings.docx"
    document = docx.Document()
    document.add_heading("Section Alpha", level=1)
    document.add_paragraph(
        "This is the first paragraph of section alpha. It contains enough text "
        "to be split into several chunks when a small token limit is applied."
    )
    document.add_paragraph(
        "Second paragraph continues the body and ensures there is measurable "
        "content for the token-limited chunker to divide."
    )
    document.add_heading("Section Beta", level=1)
    document.add_paragraph(
        "Beta section has a single but reasonably long paragraph so the heading "
        "preservation and body recovery rules can be verified independently."
    )
    document.save(str(path))
    return path


def _token_count(text: str, tokenizer: AutoTokenizer) -> int:
    return len(tokenizer(text, add_special_tokens=False)["input_ids"])


def test_chunk_document_returns_base_chunks():
    chunks = chunk_document(SAMPLE_PDF, max_tokens=150)

    assert isinstance(chunks, list)
    assert len(chunks) > 0
    for chunk in chunks:
        assert isinstance(chunk, BaseChunk)
        assert chunk.text.strip()


def test_chunk_document_respects_max_tokens(tokenizer):
    max_tokens = 150
    chunks = chunk_document(SAMPLE_PDF, max_tokens=max_tokens)

    for chunk in chunks:
        token_count = _token_count(chunk.text, tokenizer)
        assert token_count <= max_tokens, (
            f"Chunk excede el límite de tokens ({token_count} > {max_tokens}): "
            f"{chunk.text[:80]!r}..."
        )


@pytest.mark.parametrize("max_tokens", [50, 100, 250])
def test_chunk_document_with_various_limits(max_tokens: int, tokenizer):
    chunks = chunk_document(SAMPLE_PDF, max_tokens=max_tokens)

    assert len(chunks) > 0
    for chunk in chunks:
        token_count = _token_count(chunk.text, tokenizer)
        assert token_count <= max_tokens


def test_chunk_document_preserves_section_headings():
    chunks = chunk_document(SAMPLE_PDF, max_tokens=150)
    chunks_with_headings = [c for c in chunks if c.meta and c.meta.headings]

    assert len(chunks_with_headings) > 0, "Se esperaba al menos un chunk con heading"
    for chunk in chunks_with_headings:
        heading = chunk.meta.headings[0]
        assert chunk.text.startswith(heading), (
            f"Chunk no conserva el título de sección: {chunk.text[:80]!r}"
        )


def test_split_chunks_preserve_section_headings(tokenizer, sample_document):
    max_tokens = 50
    chunker = MaxTokenLimitingChunker(max_tokens=max_tokens)
    chunks = list(chunker.chunk(sample_document))

    chunks_with_headings = [c for c in chunks if c.meta and c.meta.headings]
    assert len(chunks_with_headings) > 0
    for chunk in chunks_with_headings:
        heading = chunk.meta.headings[0]
        assert chunk.text.startswith(heading)
        assert _token_count(chunk.text, tokenizer) <= max_tokens


def test_empty_document_returns_empty_list():
    empty_doc = DLDocument(name="empty")
    chunker = MaxTokenLimitingChunker(max_tokens=512)
    assert list(chunker.chunk(empty_doc)) == []


def test_max_tokens_must_be_positive():
    with pytest.raises(ValidationError):
        MaxTokenLimitingChunker(max_tokens=0)
    with pytest.raises(ValidationError):
        MaxTokenLimitingChunker(max_tokens=-1)


def test_very_small_max_tokens_does_not_crash(sample_document):
    """El chunker no debe lanzar ValueError cuando el heading ocupa todo el presupuesto."""
    max_tokens = 1
    chunker = MaxTokenLimitingChunker(max_tokens=max_tokens)
    chunks = list(chunker.chunk(sample_document))
    assert chunks, "Se esperaban chunks incluso con max_tokens=1"


def test_small_max_tokens_loses_no_section(sample_document):
    """Con max_tokens=5 ningún título ni cuerpo se pierde por no caber el heading."""
    from collections import defaultdict

    max_tokens = 5
    inner_chunks = list(HierarchicalChunker().chunk(sample_document))
    chunks = list(MaxTokenLimitingChunker(max_tokens=max_tokens).chunk(sample_document))

    assert inner_chunks, "Se esperaban chunks del chunker base"
    assert chunks, "El chunker no debe descartar todas las secciones"

    def _compact(value: str) -> str:
        return "".join(value.split())

    def _heading_tuple(chunk: BaseChunk):
        return tuple(chunk.meta.headings) if chunk.meta and chunk.meta.headings else ()

    inner_by_heading: dict[tuple[str, ...], list[BaseChunk]] = defaultdict(list)
    for c in inner_chunks:
        inner_by_heading[_heading_tuple(c)].append(c)

    out_by_heading: dict[tuple[str, ...], list[BaseChunk]] = defaultdict(list)
    for c in chunks:
        out_by_heading[_heading_tuple(c)].append(c)

    for heading, inners in inner_by_heading.items():
        assert heading in out_by_heading, (
            f"Sección perdida con max_tokens={max_tokens}: {heading}"
        )
        inner_body = _compact("".join(c.text or "" for c in inners))
        out_body = _compact("".join(c.text or "" for c in out_by_heading[heading]))
        for part in heading:
            out_body = out_body.replace(_compact(part), "")
        assert inner_body in out_body, (
            f"Se perdió texto del cuerpo con max_tokens={max_tokens}: "
            f"sección={heading}, cuerpo={inner_body[:80]!r}..."
        )


def test_chunk_document_rejects_missing_path():
    with pytest.raises((ConversionError, FileNotFoundError)):
        chunk_document("/nonexistent/file.pdf", max_tokens=150)


def test_chunk_document_accepts_string_and_path(tokenizer):
    chunks_str = chunk_document(str(SAMPLE_PDF), max_tokens=150)
    chunks_path = chunk_document(SAMPLE_PDF, max_tokens=150)

    assert len(chunks_str) == len(chunks_path)
    assert len(chunks_str) > 0
    for chunk_str, chunk_path in zip(chunks_str, chunks_path):
        assert chunk_str.text == chunk_path.text
        assert _token_count(chunk_str.text, tokenizer) <= 150


def test_max_token_limiting_chunker_class_signature():
    """El chunker expone la interfaz BaseChunker y conserva la firma."""
    chunker = MaxTokenLimitingChunker(max_tokens=150)
    assert chunker.max_tokens == 150


# ---------------------------------------------------------------------------
# DOCX contract tests (resources/rel18/rel_14.docx + synthetic heading fixture)
# ---------------------------------------------------------------------------


def test_chunk_document_docx_returns_base_chunks():
    chunks = chunk_document(SAMPLE_DOCX, max_tokens=150)

    assert isinstance(chunks, list)
    assert len(chunks) > 0
    for chunk in chunks:
        assert isinstance(chunk, BaseChunk)
        assert chunk.text.strip()


def test_chunk_document_docx_respects_max_tokens(tokenizer):
    max_tokens = 150
    chunks = chunk_document(SAMPLE_DOCX, max_tokens=max_tokens)

    for chunk in chunks:
        token_count = _token_count(chunk.text, tokenizer)
        assert token_count <= max_tokens, (
            f"Chunk excede el límite de tokens ({token_count} > {max_tokens}): "
            f"{chunk.text[:80]!r}..."
        )


@pytest.mark.parametrize("max_tokens", [50, 100, 150, 200, 250])
def test_chunk_document_docx_with_various_limits(max_tokens: int, tokenizer):
    chunks = chunk_document(SAMPLE_DOCX, max_tokens=max_tokens)

    assert len(chunks) > 0
    for chunk in chunks:
        token_count = _token_count(chunk.text, tokenizer)
        assert token_count <= max_tokens


def test_chunk_document_docx_no_word_split(tokenizer):
    """El chunker no debe partir palabras entre chunks.

    Se reconstruye el vocabulario completo del documento a partir de los
    chunks generados con distintos límites y se compara con el chunker base.
    """
    doc = DocumentConverter().convert(str(SAMPLE_DOCX)).document
    inner_chunks = list(HierarchicalChunker().chunk(doc))
    inner_words = " ".join(c.text for c in inner_chunks).split()

    for max_tokens in [50, 100, 150, 200, 250]:
        chunks = list(MaxTokenLimitingChunker(max_tokens=max_tokens).chunk(doc))
        assert len(chunks) > 0

        for chunk in chunks:
            assert _token_count(chunk.text, tokenizer) <= max_tokens

        outer_words = " ".join(c.text for c in chunks).split()
        assert outer_words == inner_words, (
            f"Palabras perdidas o partidas con max_tokens={max_tokens}: "
            f"esperadas {len(inner_words)}, obtenidas {len(outer_words)}"
        )


def test_chunk_document_docx_accepts_string_and_path(tokenizer):
    chunks_str = chunk_document(str(SAMPLE_DOCX), max_tokens=150)
    chunks_path = chunk_document(SAMPLE_DOCX, max_tokens=150)

    assert len(chunks_str) == len(chunks_path)
    assert len(chunks_str) > 0
    for chunk_str, chunk_path in zip(chunks_str, chunks_path):
        assert chunk_str.text == chunk_path.text
        assert _token_count(chunk_str.text, tokenizer) <= 150


def test_chunk_document_docx_with_headings_preserves_section_headings(
    tokenizer, docx_with_headings
):
    """DOCX con estilos de título conserva el heading como prefijo del chunk."""
    max_tokens = 150
    chunks = chunk_document(docx_with_headings, max_tokens=max_tokens)

    chunks_with_headings = [c for c in chunks if c.meta and c.meta.headings]
    assert len(chunks_with_headings) > 0, "Se esperaba al menos un chunk con heading"
    for chunk in chunks_with_headings:
        heading = chunk.meta.headings[0]
        assert chunk.text.startswith(heading), (
            f"Chunk no conserva el título de sección: {chunk.text[:80]!r}"
        )
        assert _token_count(chunk.text, tokenizer) <= max_tokens


def test_chunk_document_docx_with_headings_small_limit_keeps_body(
    tokenizer, docx_with_headings
):
    """Si el heading no cabe, se conserva en meta.headings y no se pierde cuerpo."""
    max_tokens = 5
    inner_chunks = list(
        HierarchicalChunker().chunk(
            DocumentConverter().convert(str(docx_with_headings)).document
        )
    )
    chunks = chunk_document(docx_with_headings, max_tokens=max_tokens)

    assert inner_chunks, "Se esperaban chunks del chunker base"
    assert chunks, "El chunker no debe descartar todas las secciones"

    def _compact(value: str) -> str:
        return "".join(value.split())

    def _heading_tuple(chunk: BaseChunk):
        return tuple(chunk.meta.headings) if chunk.meta and chunk.meta.headings else ()

    inner_by_heading: dict[tuple[str, ...], list[BaseChunk]] = {}
    for c in inner_chunks:
        key = _heading_tuple(c)
        inner_by_heading.setdefault(key, []).append(c)

    out_by_heading: dict[tuple[str, ...], list[BaseChunk]] = {}
    for c in chunks:
        key = _heading_tuple(c)
        out_by_heading.setdefault(key, []).append(c)

    for heading, inners in inner_by_heading.items():
        assert heading in out_by_heading, (
            f"Sección perdida con max_tokens={max_tokens}: {heading}"
        )
        inner_body = _compact("".join(c.text or "" for c in inners))
        out_body = _compact("".join(c.text or "" for c in out_by_heading[heading]))
        for part in heading:
            out_body = out_body.replace(_compact(part), "")
        assert inner_body in out_body, (
            f"Se perdió texto del cuerpo con max_tokens={max_tokens}: "
            f"sección={heading}, cuerpo={inner_body[:80]!r}..."
        )

    for chunk in chunks:
        assert _token_count(chunk.text, tokenizer) <= max_tokens


def test_chunk_document_docx_very_small_max_tokens_does_not_crash(docx_with_headings):
    """El chunker no debe lanzar ValueError con presupuestos mínimos."""
    max_tokens = 1
    chunks = chunk_document(docx_with_headings, max_tokens=max_tokens)
    assert chunks, "Se esperaban chunks incluso con max_tokens=1"
